# -*- coding: utf-8 -*-
"""
Renderizador ATS: um .docx (e, em alguns perfis, um .txt) por destino e idioma,
a partir de conteudo.py. Substitui o antigo gerar_curriculo_ptbr.py.

POR QUE UM ARQUIVO POR DESTINO
------------------------------
Nao e porque os parsers sejam muito diferentes - eles nao sao. Layout de coluna
unica, sem tabela, sem caixa de texto, titulos de secao padrao, fonte sem serifa
de 10-11pt, contato fora de cabecalho/rodape e bullets consistentes sao exigidos
por todos. A unica divergencia com evidencia e o FORMATO DE DATA (ver PERFIS).
O valor de ter um arquivo por destino esta no nome do arquivo: voce sabe o que
esta subindo. O risco normal de manter varios curriculos - subir o desatualizado
- nao existe aqui porque todos saem de conteudo.py.

POR QUE O DOCUMENTO E ASSIM
---------------------------
O alvo nao e "o texto sai do arquivo" - e "o parser preenche os campos certos".

  * Linha 1 = SO o nome. Linha 2 = e-mail cru, minusculo, sem rotulo: e a unica
    linha util sem nenhum token capitalizado, entao o parser que cola linha 1 com
    linha 2 nao tem de onde tirar um sobrenome errado.
  * Nenhum nome de tecnologia nas 6 primeiras linhas. Marca perto do nome vira
    sobrenome: o cabecalho antigo do index.html fazia o Workday da Accenture
    devolver "sobrenome: React", porque o NER marcava "Java Backend Engineer" e
    "Spring Boot" como entidade PESSOA.
  * Cargo pretendido isolado e rotulado, sem marca - e o que alimenta o campo
    "cargo".
  * Pretensao salarial longe do bloco de contato: "R$ 8.500" e alvo facil de
    regex gulosa de telefone/documento.
  * Sem travessao, meia-risca, aspa curva ou sinal de menos tipografico: tudo
    normalizado para ASCII (modelo.limpar). Acentos sao preservados.
  * Sem tabela, sem caixa de texto, sem coluna, sem cabecalho/rodape do Word.
    Titulos em estilo Heading real, bullets em lista real.

Validado em 15/09/2026 no autofill do Workday da Accenture: nome "Lucas",
sobrenome "Bueno Cesario".
"""

import re
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import conteudo as c
from modelo import T, competencias, formatar_data, formatar_periodo, limpar, rotulo

# --------------------------------------------------------------- perfis ----
#
# "datas":  "mm/aaaa" -> 11/2025 - 06/2026 | "mmm-en" -> Nov 2025 - Jun 2026
#           "mmm-pt"  -> Nov 2025 - Jun 2026 com Ago/Abr/Out/Fev/Set/Dez
# A diferenca de data e a unica com evidencia. O Greenhouse e documentado como
# exigente com "MMM YYYY"; Gupy e LinkedIn sao nativos em mm/aaaa.

PERFIS = {
    "Workday": {
        "datas": "mmm-en", "salario": False, "txt": True, "idiomas": ("pt", "en"),
        "nota": "Nome validado na Accenture em 15/09/2026; experiencia NAO foi "
                "extraida com datas em mm/aaaa. A documentacao do Workday usa "
                "'Jun 2022 - Present', entao aqui e mmm-en. PENDENTE DE RETESTE.",
    },
    # Candidata B do reteste: identica a de Workday, so muda a abreviacao do mes
    # para portugues. Apagar a perdedora depois do reteste.
    "Workday-mesPT": {
        "datas": "mmm-pt", "salario": False, "txt": False, "idiomas": ("pt",),
        "nota": "Candidata B do reteste no Workday. PENDENTE.",
    },
    "Gupy": {
        "datas": "mm/aaaa", "salario": False, "txt": False, "idiomas": ("pt",),
        "nota": "DOCX e mais confiavel que PDF no parser Gaia.",
    },
    "LinkedIn": {
        "datas": "mm/aaaa", "salario": False, "txt": False, "idiomas": ("pt", "en"),
        "nota": "Sem pretensao: o LinkedIn importa para o perfil, que e publico.",
    },
    "Greenhouse": {
        "datas": "mmm-en", "salario": False, "txt": False, "idiomas": ("pt", "en"),
        "nota": "Datas em MMM YYYY. Greenhouse e majoritariamente empresa "
                "internacional - na duvida, mande a versao em ingles.",
    },
    "Analise": {
        "datas": "mmm-en", "salario": False, "txt": True, "idiomas": ("pt", "en"),
        "nota": "Para Jobscan, Resume Worded e afins. Sem pretensao, para nao "
                "poluir a contagem de keyword.",
    },
}

SECOES = {
    "resumo": T("Resumo profissional", "Professional summary"),
    "competencias": T("Competências técnicas", "Technical skills"),
    "experiencia": T("Experiência profissional", "Professional experience"),
    "formacao": T("Formação acadêmica", "Education"),
    "certificacoes": T("Certificações", "Certifications"),
    "idiomas": T("Idiomas", "Languages"),
    "adicionais": T("Informações adicionais", "Additional information"),
}

ROT_TECNOLOGIAS = T("Tecnologias", "Technologies")
ROT_CONCLUSAO = T("Conclusão prevista", "Expected completion")

# O nome do arquivo e sinal de verdade para varios parsers: e a primeira coisa
# que eles usam para conferir o nome extraido do corpo. Nome completo primeiro.
BASE = {"pt": "Lucas-Bueno-Cesario-Curriculo", "en": "Lucas-Bueno-Cesario-Resume"}


def nome_arquivo(perfil, idioma, ext):
    return "{}-{}.{}".format(BASE[idioma], perfil, ext)


# --------------------------------------------------------------- blocos ----

def blocos(perfil, idioma):
    """Sequencia neutra do documento: [(tipo, valor), ...]. DOCX e TXT so mudam
    a forma de desenhar; a ordem e o conteudo sao decididos aqui, uma vez.

    tipos: nome, linha, rotulado (rot, val), secao, paragrafo, cargo, data,
           bullet, tecnologias (rot, val)
    """
    cfg = PERFIS[perfil]
    if idioma not in cfg["idiomas"]:
        raise ValueError("perfil {} nao tem versao em {}".format(perfil, idioma))

    b = [("nome", c.NOME), ("linha", c.EMAIL)]
    for rot, val in c.CONTATO_ATS:
        b.append(("rotulado", (rot.em(idioma), val.em(idioma))))

    b.append(("secao", SECOES["resumo"].em(idioma)))
    b.append(("paragrafo", c.RESUMO_ATS.em(idioma)))
    b.append(("paragrafo", c.DOMINIOS.em(idioma)))

    b.append(("secao", SECOES["competencias"].em(idioma)))
    for categoria, itens in competencias(c.EXPERIENCIAS, c.CATEGORIAS, idioma):
        b.append(("rotulado", (categoria, itens)))

    b.append(("secao", SECOES["experiencia"].em(idioma)))
    for vaga in c.EXPERIENCIAS:
        # Quatro linhas, uma informacao por linha, nesta ordem. A data fica
        # sozinha e imediatamente acima dos bullets: e a ancora que o parser usa
        # para fechar o registro da vaga.
        b.append(("cargo", vaga.cargo.em(idioma)))
        b.append(("linha", vaga.empresa.em(idioma)))
        b.append(("linha", vaga.meta_ats.em(idioma)))
        b.append(("data", formatar_periodo(vaga.inicio, vaga.fim, cfg["datas"])))
        for bullet in vaga.bullets:
            if bullet.ats is not None:
                b.append(("bullet", bullet.ats.em(idioma)))
        tecnologias = ", ".join(rotulo(k, idioma, c.ROTULOS) for k in vaga.tecnologias)
        b.append(("tecnologias", (ROT_TECNOLOGIAS.em(idioma), tecnologias)))

    b.append(("secao", SECOES["formacao"].em(idioma)))
    for formacao in c.FORMACAO:
        b.append(("cargo", formacao.curso.em(idioma)))
        b.append(("linha", formacao.instituicao.em(idioma)))
        b.append(("linha", "{}: {}".format(ROT_CONCLUSAO.em(idioma),
                                          formatar_data(formacao.conclusao, cfg["datas"]))))

    b.append(("secao", SECOES["certificacoes"].em(idioma)))
    for certificacao in c.CERTIFICACOES:
        b.append(("bullet", certificacao.em(idioma)))

    b.append(("secao", SECOES["idiomas"].em(idioma)))
    for nome, nivel in c.IDIOMAS_ATS:
        b.append(("rotulado", (nome.em(idioma), nivel.em(idioma))))

    b.append(("secao", SECOES["adicionais"].em(idioma)))
    adicionais = [c.PRETENSAO, c.DISPONIBILIDADE] if cfg["salario"] else [c.DISPONIBILIDADE]
    for rot, val in adicionais:
        b.append(("rotulado", (rot.em(idioma), val.em(idioma))))
    return b


# ---------------------------------------------------------------- docx -----

FONTE = "Arial"
PRETO = RGBColor(0, 0, 0)


def _forcar_fonte(run):
    """Fixa a fonte tambem para complex-script/east-asian, senao o Word troca a
    fonte de trechos acentuados em alguns ambientes."""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONTE)


def _estilar(run, tamanho, negrito=False):
    run.font.name = FONTE
    run.font.size = Pt(tamanho)
    run.font.bold = negrito
    run.font.color.rgb = PRETO
    _forcar_fonte(run)
    return run


def _p(doc, texto, tamanho=10, negrito=False, espaco_antes=0, espaco_depois=2):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(espaco_antes)
    par.paragraph_format.space_after = Pt(espaco_depois)
    par.paragraph_format.line_spacing = 1.08
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _estilar(par.add_run(limpar(texto)), tamanho, negrito)
    return par


def _p_rotulado(doc, rotulo_, valor, tamanho=10, espaco_depois=2):
    """Paragrafo unico 'Rotulo: valor'. Um paragrafo so garante que a extracao
    devolva o rotulo e o valor na mesma linha."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(espaco_depois)
    par.paragraph_format.line_spacing = 1.08
    _estilar(par.add_run(limpar(rotulo_) + ": "), tamanho, negrito=True)
    _estilar(par.add_run(limpar(valor)), tamanho)
    return par


def _secao(doc, titulo):
    """Titulo de secao em estilo 'Heading 1' real: o nome do estilo e o que
    muitos parsers usam para delimitar as secoes."""
    par = doc.add_paragraph(style="Heading 1")
    par.paragraph_format.space_before = Pt(12)
    par.paragraph_format.space_after = Pt(4)
    _estilar(par.add_run(limpar(titulo).upper()), 11.5, negrito=True)

    ppr = par._p.get_or_add_pPr()
    pbdr = ppr.makeelement(qn("w:pBdr"), {})
    bottom = pbdr.makeelement(qn("w:bottom"), {})
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)
    ppr.append(pbdr)
    return par


def _bullet(doc, texto):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.08
    par.paragraph_format.left_indent = Cm(0.63)
    _estilar(par.add_run(limpar(texto)), 10)
    return par


def gerar_docx(caminho, perfil, idioma):
    doc = Document()

    secao = doc.sections[0]
    secao.top_margin = Cm(1.6)
    secao.bottom_margin = Cm(1.6)
    secao.left_margin = Cm(1.9)
    secao.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(10)
    normal.font.color.rgb = PRETO
    rfonts = normal.element.rPr.rFonts
    rfonts.set(qn("w:eastAsia"), FONTE)
    rfonts.set(qn("w:cs"), FONTE)

    # Cabecalho como paragrafo normal, nunca no header do Word - boa parte dos
    # ATS descarta o conteudo de cabecalho/rodape.
    anterior = None
    for tipo, valor in blocos(perfil, idioma):
        if tipo == "nome":
            _p(doc, valor, tamanho=20, negrito=True, espaco_depois=4)
        elif tipo == "linha":
            _p(doc, valor, espaco_depois=1)
        elif tipo == "data":
            _p(doc, valor, espaco_depois=3)
        elif tipo == "rotulado":
            _p_rotulado(doc, valor[0], valor[1], espaco_depois=1)
        elif tipo == "tecnologias":
            _p_rotulado(doc, valor[0], valor[1], tamanho=9.5)
        elif tipo == "secao":
            _secao(doc, valor)
        elif tipo == "paragrafo":
            _p(doc, valor, espaco_depois=4)
        elif tipo == "cargo":
            # O cargo NAO usa estilo Heading, so negrito. Estilo de titulo e o que
            # o parser usa para delimitar SECAO; com cada vaga em Heading 2 o bloco
            # de experiencia se fragmenta e o Workday nao monta o historico. Foi
            # assim que a extracao de experiencia falhou na Accenture.
            _p(doc, valor, tamanho=11, negrito=True,
               espaco_antes=4 if anterior == "secao" else 8, espaco_depois=0)
        elif tipo == "bullet":
            _bullet(doc, valor)
        else:
            raise ValueError("tipo de bloco desconhecido: " + tipo)
        anterior = tipo

    # Metadados: varios parsers leem dc:creator / dc:title antes do corpo.
    props = doc.core_properties
    props.author = c.NOME
    props.last_modified_by = c.NOME
    props.title = c.NOME + (" - Curriculo" if idioma == "pt" else " - Resume")
    props.subject = c.CARGO_ALVO.em(idioma)
    props.category = c.CARGO_ALVO.em(idioma)
    props.language = "pt-BR" if idioma == "pt" else "en-US"
    props.keywords = c.PALAVRAS_CHAVE.em(idioma)[:255]

    caminho = Path(caminho)
    doc.save(str(caminho))
    return caminho


# ----------------------------------------------------------------- pdf -----

def gerar_pdf(caminho, perfil, idioma):
    """PDF do mesmo conteudo do DOCX (mesma ordem, coluna unica, sem tabela),
    impresso pelo Chrome. Texto real, nao imagem: o parser continua lendo."""
    import humano  # so aqui: humano nao precisa do ats para o resto

    corpo = []
    em_lista = False
    for tipo, valor in blocos(perfil, idioma):
        if tipo != "bullet" and em_lista:
            corpo.append("</ul>")
            em_lista = False
        if tipo == "bullet":
            if not em_lista:
                corpo.append("<ul>")
                em_lista = True
            corpo.append("<li>{}</li>".format(escape(limpar(valor))))
        elif tipo == "nome":
            corpo.append("<h1>{}</h1>".format(escape(limpar(valor))))
        elif tipo == "secao":
            corpo.append("<h2>{}</h2>".format(escape(limpar(valor).upper())))
        elif tipo == "cargo":
            corpo.append('<p class="cargo">{}</p>'.format(escape(limpar(valor))))
        elif tipo in ("rotulado", "tecnologias"):
            corpo.append("<p><b>{}:</b> {}</p>".format(escape(limpar(valor[0])),
                                                        escape(limpar(valor[1]))))
        else:  # linha, data, paragrafo
            corpo.append("<p>{}</p>".format(escape(limpar(valor))))
    if em_lista:
        corpo.append("</ul>")

    html = (
        '<!doctype html><html lang="{lang}"><head><meta charset="utf-8">'
        "<title>{titulo}</title><style>"
        "@page{{size:Letter;margin:1.6cm 1.9cm}}"
        "body{{font:10pt/1.25 Arial,sans-serif;color:#000;margin:0}}"
        "h1{{font-size:20pt;margin:0 0 4pt}}"
        "h2{{font-size:11.5pt;margin:12pt 0 4pt;border-bottom:1px solid #000;"
        "padding-bottom:1pt;break-after:avoid}}"
        "p{{margin:0 0 2pt}} p.cargo{{font-size:11pt;font-weight:bold;margin-top:8pt;"
        "break-after:avoid}}"
        "ul{{margin:0 0 2pt;padding-left:0.63cm}} li{{margin-bottom:2pt}}"
        "</style></head><body>{corpo}</body></html>"
    ).format(lang="pt-BR" if idioma == "pt" else "en-US",
             titulo=escape(c.NOME + (" - Curriculo" if idioma == "pt" else " - Resume")),
             corpo="\n".join(corpo))
    return humano.imprimir_pdf(html, caminho)


# ----------------------------------------------------------------- txt -----

def gerar_txt(caminho, perfil, idioma):
    linhas = []
    for tipo, valor in blocos(perfil, idioma):
        if tipo in ("nome", "linha", "data"):
            linhas.append(valor)
        elif tipo in ("rotulado", "tecnologias"):
            linhas.append(valor[0] + ": " + valor[1])
        elif tipo == "secao":
            linhas.extend(["", valor.upper(), "=" * len(valor)])
        elif tipo == "paragrafo":
            linhas.extend([valor, ""])
        elif tipo == "cargo":
            linhas.extend(["", valor])
        elif tipo == "bullet":
            linhas.append("- " + valor)
        else:
            raise ValueError("tipo de bloco desconhecido: " + tipo)
    texto = re.sub(r"\n{3,}", "\n\n", "\n".join(linhas))
    caminho = Path(caminho)
    caminho.write_text(limpar(texto).strip() + "\n", encoding="utf-8")
    return caminho
