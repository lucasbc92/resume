import re

import pytest
from docx import Document

import ats
import conteudo as c


def _pars(caminho):
    return [p.text for p in Document(str(caminho)).paragraphs]


def _docx(tmp_path, perfil, idioma):
    return ats.gerar_docx(tmp_path / ats.nome_arquivo(perfil, idioma, "docx"), perfil, idioma)


def _txt(tmp_path, perfil, idioma):
    caminho = ats.gerar_txt(tmp_path / ats.nome_arquivo(perfil, idioma, "txt"), perfil, idioma)
    return caminho.read_text(encoding="utf-8")


COMBINACOES = [(p, i) for p, cfg in ats.PERFIS.items() for i in cfg["idiomas"]]


def test_nomes_de_arquivo():
    assert ats.nome_arquivo("Gupy", "pt", "docx") == "Lucas-Bueno-Cesario-Curriculo-Gupy.docx"
    assert ats.nome_arquivo("Workday", "en", "txt") == "Lucas-Bueno-Cesario-Resume-Workday.txt"


@pytest.mark.parametrize("perfil, idioma", COMBINACOES)
def test_linha_1_nome_linha_2_email(tmp_path, perfil, idioma):
    pars = _pars(_docx(tmp_path, perfil, idioma))
    assert pars[0] == c.NOME
    assert pars[1] == c.EMAIL


@pytest.mark.parametrize("perfil, idioma", COMBINACOES)
def test_nenhuma_tecnologia_nas_6_primeiras_linhas(tmp_path, perfil, idioma):
    chaves = {k for _, entradas in c.CATEGORIAS for _, ks in entradas for k in ks}
    for linha in _pars(_docx(tmp_path, perfil, idioma))[:6]:
        for chave in chaves:
            padrao = r"(?<!\w)" + re.escape(chave) + r"(?!\w)"
            assert not re.search(padrao, linha), (chave, linha)


def test_secoes_sao_heading_1_e_cargos_nao(tmp_path):
    doc = Document(str(_docx(tmp_path, "Workday", "pt")))
    estilos = {p.text: p.style.name for p in doc.paragraphs}
    assert estilos["EXPERIÊNCIA PROFISSIONAL"] == "Heading 1"
    assert estilos["Desenvolvedor Java Full Stack Sênior"] == "Normal"


@pytest.mark.parametrize("perfil, idioma, esperado", [
    ("Gupy", "pt", "08/2025 - 11/2025"),
    ("Workday-mesPT", "pt", "Ago 2025 - Nov 2025"),
    ("Workday", "en", "Aug 2025 - Nov 2025"),
])
def test_datas_no_formato_do_perfil(tmp_path, perfil, idioma, esperado):
    assert esperado in _pars(_docx(tmp_path, perfil, idioma))


@pytest.mark.parametrize("perfil, idioma, tem", [
    ("Workday", "pt", True), ("Gupy", "pt", True), ("LinkedIn", "pt", False),
    ("Analise", "pt", False), ("Workday", "en", True), ("LinkedIn", "en", False),
])
def test_pretensao_so_nos_perfis_com_salario(tmp_path, perfil, idioma, tem):
    texto = "\n".join(_pars(_docx(tmp_path, perfil, idioma)))
    marca = "R$ 8.500" if idioma == "pt" else "US$ 25-35"
    assert (marca in texto) is tem
    if not tem:
        assert "R$" not in texto and "US$" not in texto


def test_pontuacao_ascii_e_acentos_preservados(tmp_path):
    texto = _txt(tmp_path, "Workday", "pt")
    for simbolo in "\u2014\u2013\u2212\u201c\u201d\u2018\u2019\u00b7\u00a0":
        assert simbolo not in texto, repr(simbolo)
    assert "Sênior" in texto and "Telefónica" in texto


def test_ingles_sem_rotulos_em_portugues(tmp_path):
    texto = _txt(tmp_path, "Workday", "en")
    for pt in ("Tecnologias", "Experiência", "Conclusão", "Pretensão", "Competências",
               "Remoto", "Ago 20", "Abr 20"):
        assert pt not in texto, pt
    assert "PROFESSIONAL EXPERIENCE" in texto and "Technologies:" in texto


def test_txt_tem_secoes_sublinhadas(tmp_path):
    linhas = _txt(tmp_path, "Analise", "pt").splitlines()
    i = linhas.index("EXPERIÊNCIA PROFISSIONAL")
    assert linhas[i + 1] == "=" * len("Experiência profissional")


def test_bullet_so_humano_nao_entra_no_ats(tmp_path):
    assert "rotinas de backup em Linux." not in _txt(tmp_path, "Workday", "pt")


def test_perfil_so_em_portugues_rejeita_ingles():
    with pytest.raises(ValueError):
        ats.blocos("Gupy", "en")


def test_metadados_do_docx(tmp_path):
    props = Document(str(_docx(tmp_path, "Workday", "en"))).core_properties
    assert props.author == c.NOME
    assert props.subject == c.CARGO_ALVO.en


def test_docx_sem_pontuacao_tipografica(tmp_path):
    texto = "\n".join(_pars(_docx(tmp_path, "Workday", "pt")))
    for simbolo in "\u2014\u2013\u2212\u201c\u201d\u2018\u2019\u00b7\u00a0":
        assert simbolo not in texto, repr(simbolo)
    assert "Sênior" in texto and "Telefónica" in texto


def test_contato_em_portugues_com_acentos(tmp_path):
    pars = _pars(_docx(tmp_path, "Workday", "pt"))
    assert "País: Brasil" in pars
    assert any(p.startswith("Modelo de trabalho: Remoto (fuso UTC-3; sobreposição total")
               for p in pars)
